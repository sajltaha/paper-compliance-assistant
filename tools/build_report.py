from pathlib import Path
import json,argparse,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ap=argparse.ArgumentParser()
ap.add_argument('--source-sha',required=True)
ap.add_argument('--output-dir',required=True)
ap.add_argument('--local-amendments',default='')
args=ap.parse_args()
if not re.fullmatch(r'[0-9a-f]{40}',args.source_sha):raise SystemExit('Use the exact immutable 40-character source commit')
ROOT=Path(__file__).resolve().parents[1]
OUT=Path(args.output_dir);OUT.mkdir(parents=True,exist_ok=True)
sections=json.loads((ROOT/'docs/report_content.json').read_text().replace('{{REVISION}}',args.source_sha))
def make_doc():
 d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
 styles=d.styles
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
  styles[name].font.name='DejaVu Sans';styles[name].font.color.rgb=RGBColor(0,0,0)
 styles['Normal'].font.size=Pt(10);styles['Normal'].paragraph_format.space_after=Pt(7);styles['Normal'].paragraph_format.line_spacing=1.12
 styles['Title'].font.size=Pt(23);styles['Title'].font.bold=True
 styles['Heading 1'].font.size=Pt(15);styles['Heading 1'].paragraph_format.space_before=Pt(15);styles['Heading 1'].paragraph_format.space_after=Pt(9)
 styles['Heading 2'].font.size=Pt(11);styles['Heading 2'].paragraph_format.space_before=Pt(9)
 reference_style=styles.add_style('Reference text',1);reference_style.base_style=styles['Normal'];reference_style.font.size=Pt(9);reference_style.paragraph_format.space_after=Pt(5)
 for st in styles:
  pr=st.element.find(qn('w:pPr'))
  if pr is not None:
   for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
 header=sec.header.paragraphs[0];header.text='AITU  |  SOFTWARE ENGINEERING  |  FIRST PROBATIONARY PERIOD';header.style='Caption';header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor(0,0,0)
 foot=sec.footer.paragraphs[0];foot.alignment=2
 foot.add_run('Checkpoint report 24 October 2026  |  ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 for r in foot.runs:r.font.size=Pt(8)
 d.add_paragraph('Scientific Paper Compliance Assistant','Title')
 d.add_paragraph('First probationary period report','Subtitle')
 d.add_paragraph('Prepared for supervisor and committee review\nSaylaubekov Talgat Yerkhanuly  |  SE-2429\nNuryshev Kaisar Serikuly  |  SE-2413\nSupervisor  Soltan Gulzhan Zheksenbayevna\n24 October 2026')
 p=d.add_paragraph();r=p.add_run('Status  ');r.bold=True;p.add_run('Research and design prepared. Supervisor response received on 24 October at 03:00; no revisions requested. Remaining personal acceptance and final submission are tracked in YouTrack. Source revision ' + args.source_sha[:12] + ('; local amendments ' + args.local_amendments[:12] if args.local_amendments else '') + '.')
 d.add_paragraph('Contents','Heading 2')
 d.add_paragraph('1 Proposal   2 Relevance   3 Literature   4 Existing systems   5 Technologies\n6 Requirements   7 Architecture   8 UX prototype   9 Evaluation\n10 Management and risks   11 Supervisor feedback   References')
 for si,(title,blocks) in enumerate(sections):
  d.add_heading(title,level=1)
  for b in blocks:
   if b[0]=='p':
    paragraph=d.add_paragraph(b[1])
    if title=='References':
     paragraph.style='Reference text'
   elif b[0]=='h':d.add_heading(b[1],level=2)
   elif b[0]=='image':
    p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(ROOT/'diagrams'/f'{b[1]}.png'),width=Inches(6.65))
    cap=d.add_paragraph(b[2]);cap.style='Caption';cap.runs[0].font.color.rgb=RGBColor(0,0,0)
   elif b[0]=='table':
    heads,rows,widths=b[1:];t=d.add_table(rows=1,cols=len(heads));t.autofit=False
    if widths:
     for c,w in zip(t.columns,widths):c.width=Inches(w)
    for j,h in enumerate(heads):t.rows[0].cells[j].text=h
    trpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trpr.append(rep)
    for row in rows:
     c=t.add_row().cells
     for j,txt in enumerate(row):c[j].text=str(txt)
    for ri,row in enumerate(t.rows):
     trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
     for j,c in enumerate(row.cells):
      if widths:c.width=Inches(widths[j])
      c.vertical_alignment=1
      pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
      for edge in ['top','left','bottom','right']:
       e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
      pr.append(bd);m=OxmlElement('w:tcMar')
      for edge in ['top','left','bottom','right']:
       e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');m.append(e)
      pr.append(m)
      sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'243C50' if ri==0 else ('F1F5F7' if ri%2==0 else 'FFFFFF'));pr.append(sh)
      for p in c.paragraphs:
       p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.07
       for r in p.runs:r.font.size=Pt(8 if len(heads)>3 else 8.6);r.font.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
 d.save(OUT/'SE-2429_Saylaubekov_ProbPeriod1.docx')

make_doc()
p=OUT/'SE-2429_Saylaubekov_ProbPeriod1.docx'
d=Document(p)
for para in d.paragraphs:
 text=para.text
 if 'https://' not in text:continue
 para.clear()
 for part in re.split(r'(https://[^\s]+)',text):
  if not part.startswith('https://'):para.add_run(part);continue
  url=part.rstrip('.,;')
  h=OxmlElement('w:hyperlink');h.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
  r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);para._p.append(h)
  if part[len(url):]:para.add_run(part[len(url):])
d.save(p)
print(p)

